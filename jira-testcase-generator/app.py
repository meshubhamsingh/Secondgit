import os
import json
import shutil
import google.generativeai as genai
from fastapi import FastAPI, Request, Form, UploadFile, File
from fastapi.responses import HTMLResponse, FileResponse
from fastapi.templating import Jinja2Templates
from dotenv import load_dotenv
import tempfile
from pypdf import PdfReader
import io
from openpyxl import load_workbook

# --- 1. CONFIGURATION ---
load_dotenv()
api_key = os.getenv("GEMINI_API_KEY")

if not api_key:
    raise ValueError("API Key is missing. Check your .env file.")

# UPDATE THIS IF NEEDED (e.g., 'gemini-1.5-flash' or 'gemini-1.0-pro')
genai.configure(api_key=api_key)
MODEL_NAME = 'gemini-2.5-flash' 

app = FastAPI()
templates = Jinja2Templates(directory="templates")

TEMPLATE_FILE = "template.xlsx"

# --- 2. UTILS ---
def get_webui_context():
    try:
        with open("webui.txt", "r") as f:
            return f.read()
    except FileNotFoundError:
        return "UI Structure not found."

def extract_text_from_pdf(file_bytes):
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    except Exception as e:
        return f"Error reading PDF: {str(e)}"

def get_template_headers():
    """Reads the headers from the local template.xlsx file"""
    if not os.path.exists(TEMPLATE_FILE):
        raise FileNotFoundError(f"Please create a '{TEMPLATE_FILE}' with headers in the project folder.")
    
    wb = load_workbook(TEMPLATE_FILE)
    ws = wb.active
    # Get values from the first row
    headers = [cell.value for cell in ws[1] if cell.value is not None]
    return headers

# --- 3. GEMINI LOGIC ---
def generate_dynamic_test_cases(ac_text, desc_text, pdf_content, headers):
    ui_context = get_webui_context()
    
    # Construct Requirements
    user_reqs = ""
    if desc_text: user_reqs += f"Description: {desc_text}\n"
    if ac_text: user_reqs += f"Acceptance Criteria: {ac_text}\n"
    if pdf_content: user_reqs += f"PDF Context: {pdf_content}\n"

    # Dynamic Prompt
    prompt = f"""
    Role: You are a QA Lead.
    Task: Write manual test scripts based on the requirements, mapping them EXACTLY to the provided Excel headers.
    
    Context (Web UI):
    {ui_context}
    
    User Requirements:
    {user_reqs}
    
    Target Excel Headers:
    {json.dumps(headers)}
    
    Instructions:
    1. Output a JSON object with a key "rows".
    2. "rows" must be a list of objects, where each object represents ONE ROW in Excel.
    3. The keys in every object MUST MATCH the "Target Excel Headers" exactly.
    4. For multi-step test cases, you must repeat the common data (like Scenario Name, Script ID) in every row/object.
    5. Do not invent new keys.
    
    Example Structure (based on your headers):
    {{
      "rows": [
         {{ "{headers[0]}": "TC01", "{headers[1]}": "Verify Login", ..., "Step Action": "Enter User", "Expected Result": "..." }},
         {{ "{headers[0]}": "TC01", "{headers[1]}": "Verify Login", ..., "Step Action": "Click Login", "Expected Result": "..." }}
      ]
    }}
    
    Return ONLY valid JSON.
    """

    model = genai.GenerativeModel(MODEL_NAME)
    response = model.generate_content(prompt)
    clean_text = response.text.replace("```json", "").replace("```", "").strip()
    return json.loads(clean_text)

# --- 4. ROUTES ---
@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})

@app.post("/generate")
async def generate_excel(
    acceptance_criteria: str = Form(None), 
    description: str = Form(None),
    pdf_file: UploadFile = File(None)
):
    try:
        # 1. Read Inputs
        pdf_text = ""
        if pdf_file and pdf_file.filename:
            content = await pdf_file.read()
            pdf_text = extract_text_from_pdf(content)

        if not acceptance_criteria and not description and not pdf_text:
            return {"error": "Please provide requirements (Text or PDF)."}

        # 2. Read Template Headers
        try:
            headers = get_template_headers()
        except Exception as e:
            return {"error": str(e)}

        # 3. Get Data from Gemini
        data = generate_dynamic_test_cases(acceptance_criteria, description, pdf_text, headers)
        rows = data.get("rows", [])

        if not rows:
            return {"error": "AI generated empty data."}

        # 4. Fill Data into Template Copy
        # We create a temporary copy of the template so we don't overwrite the original
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            shutil.copy(TEMPLATE_FILE, tmp.name) # Copy original template to temp file
            
            wb = load_workbook(tmp.name)
            ws = wb.active
            
            # Find the next empty row (usually row 2)
            start_row = ws.max_row + 1
            
            for row_data in rows:
                # Create a list of values in the exact order of headers
                row_values = [row_data.get(h, "") for h in headers]
                ws.append(row_values)

            wb.save(tmp.name)
            output_path = tmp.name

        # 5. Download
        return FileResponse(
            output_path, 
            filename="Generated_Test_Scripts.xlsx",
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    except Exception as e:
        return {"error": f"An error occurred: {str(e)}"}

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)